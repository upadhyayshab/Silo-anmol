"""
Parses reference_docs/CR V2 - Scoring_Reference_v4_.xlsx into plain Python
structures. Shared by scripts/generate_seed_sql.py (Excel -> reviewable SQL file)
and tests/test_scoring_engine.py (Excel -> local throwaway DB for engine tests).
"""
from pathlib import Path

import openpyxl

WORKBOOK_PATH = Path(__file__).resolve().parents[2] / "reference_docs" / "CR V2 - Scoring_Reference_v4_.xlsx"

SEASON_CODES = {"Kharif": "KHARIF", "Rabi": "RABI", "Summer": "ZAID", "Long-term": "LONGTERM"}
WATER_CODES = {"Rainfed": "RAINFED", "Surface": "SURFACE", "Drip": "DRIP", "Sprinkler": "SPRINKLER"}


def crop_uid(code: int) -> str:
    return f"crop-{int(code):03d}"


def load_workbook():
    return openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)


def rows_after_header(ws, header_predicate):
    """Yield data rows starting right after the row that matches header_predicate."""
    started = False
    for row in ws.iter_rows(values_only=True):
        if not started:
            if row and header_predicate(row):
                started = True
            continue
        if row[0] is None:
            continue
        yield row


def parse_crops(wb) -> dict[int, dict]:
    ws = wb["2. Crop Attributes"]
    crops = {}
    for row in rows_after_header(ws, lambda r: r[0] == "#"):
        code, name, family, horizon, water_need, cost_class, main_zones, note = row[:8]
        crops[int(code)] = {
            "code": int(code),
            "name": name,
            "family": family,
            "horizon": horizon.upper(),
            "water_need": water_need.upper(),
            "cost_class": cost_class.upper(),
            "main_zones": main_zones,
            "note": note,
        }
    return crops


def parse_zone_scores(wb) -> list[tuple[int, int, float]]:
    ws = wb["3. Zone (HARD)"]
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "#"):
        code = int(row[0])
        for col_idx, zone_num in enumerate(range(1, 16), start=2):
            score = row[col_idx]
            if score is None:
                continue
            out.append((code, zone_num, score))
    return out


def parse_season_scores(wb) -> list[tuple[int, str, float]]:
    ws = wb["4. Season (HARD)"]
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "#"):
        code = int(row[0])
        # columns: #, Crop, Horizon, Kharif, Rabi, Summer, Long-term
        for season_label, col_idx in (("Kharif", 3), ("Rabi", 4), ("Summer", 5), ("Long-term", 6)):
            score = row[col_idx]
            if score is None:
                continue
            out.append((code, SEASON_CODES[season_label], score))
    return out


def parse_water_scores(wb) -> list[tuple[int, str, float]]:
    ws = wb["5. Water (HARD)"]
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "#"):
        code = int(row[0])
        # columns: #, Crop, Water need, Rainfed, Surface, Drip, Sprinkler
        for method_label, col_idx in (("Rainfed", 3), ("Surface", 4), ("Drip", 5), ("Sprinkler", 6)):
            score = row[col_idx]
            if score is None:
                continue
            out.append((code, WATER_CODES[method_label], score))
    return out


def parse_soil_scores(wb) -> list[tuple[int, int, float]]:
    ws = wb["6. Soil (score)"]
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "#"):
        code = int(row[0])
        for soil_class in range(1, 13):
            score = row[1 + soil_class]  # column 2 = class 1, ... column 13 = class 12
            if score is None:
                continue
            out.append((code, soil_class, score))
    return out


def parse_budget_rules(wb) -> list[tuple[str, str, float]]:
    ws = wb["8. Budget Logic"]
    out = []
    for row in rows_after_header(ws, lambda r: r and r[0] and str(r[0]).startswith("Farmer spend")):
        label = row[0]
        if not label or not str(label).startswith(("B1", "B2", "B3", "B4")):
            break
        bucket = str(label).split()[0]
        for cost_class, col_idx in (("LOW", 1), ("MEDIUM", 2), ("HIGH", 3)):
            score = row[col_idx]
            if score is None:
                continue
            out.append((bucket, cost_class, score))
    return out


def parse_rotation_rules(wb) -> list[tuple[str, float]]:
    ws = wb["10. Rotation Logic"]
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "Situation (previous crop FAMILY, from Q3/Q3a)"):
        label, score = row[0], row[1]
        if label is None or score is None:
            continue
        if "DIFFERENT" in label.upper():
            situation = "DIFFERENT_FAMILY"
        elif "FALLOW" in label.upper():
            situation = "FALLOW"
        elif "SAME" in label.upper():
            situation = "SAME_FAMILY"
        else:
            continue
        out.append((situation, score))
    return out


def parse_weights(wb) -> list[tuple[str, float]]:
    ws = wb["1. Weights"]
    label_to_factor = {
        "Zone": "ZONE", "Season": "SEASON", "Water-method": "WATER",
        "Soil": "SOIL", "Budget": "BUDGET", "Rotation": "ROTATION",
    }
    out = []
    for row in rows_after_header(ws, lambda r: r[0] == "Factor"):
        label, weight = row[0], row[1]
        if label not in label_to_factor:
            continue
        out.append((label_to_factor[label], weight))
    return out


__all__ = [
    "WORKBOOK_PATH", "crop_uid", "load_workbook",
    "parse_crops", "parse_zone_scores", "parse_season_scores", "parse_water_scores",
    "parse_soil_scores", "parse_budget_rules", "parse_rotation_rules", "parse_weights",
]
